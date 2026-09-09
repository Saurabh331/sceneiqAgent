from docx import Document

def create_document():
    doc = Document()
    doc.add_heading('SceneIQ: End-to-End Architecture & Description', 0)

    doc.add_heading('Introduction', level=1)
    doc.add_paragraph('SceneIQ is an agentic screenplay intelligence platform designed to convert screenplays and production documents into searchable, grounded, and actionable production knowledge. It features specialized toolsets for Filmmakers, Writers, and Film Enthusiasts.')

    doc.add_heading('End-to-End Architecture Flow', level=1)
    doc.add_paragraph('The application leverages a modern AI architecture combining Retrieval-Augmented Generation (RAG) with real-world tool execution (MCP).')

    doc.add_heading('1. Streamlit UI (Frontend Layer)', level=2)
    doc.add_paragraph('- Users interact with the platform through a Streamlit web interface.\n- It provides a chat interface and dedicated tabs for specialized tools (Producers, Writers, Enthusiasts).\n- Sends requests containing user queries and (optional) session IDs to the FastAPI backend.')

    doc.add_heading('2. FastAPI Backend (Orchestration Layer)', level=2)
    doc.add_paragraph('- Serves as the central router connecting the UI to various logic modules.\n- Evaluates which tool or agent loop to invoke based on the requested endpoint.')

    doc.add_heading('3. Context Retrieval (Data Layer)', level=2)
    doc.add_paragraph('Before sending prompts to the LLM, the backend gathers relevant context from two primary sources:\n- BigQuery Vector Search: Natively queries stored script embeddings in BigQuery to find exact scenes, characters, and plot points.\n- Parallel Search MCP: Uses the Model Context Protocol to execute web searches for real-world data not found in the script.')

    doc.add_heading('4. Google Gemini (Reasoning & Generation Layer)', level=2)
    doc.add_paragraph('- Gemini 2.5 Pro: Deep creative reasoning, plot synthesis, and complex analysis.\n- Gemini 2.5 Flash: High-speed lookups, structured schema extractions, and rapid chat responses.\n- Gemini Visual/Imagen: Handles multimodal tasks like cross-referencing audition videos or generating storyboard concepts.')

    doc.add_heading('5. Downstream Applications (Output Layer)', level=2)
    doc.add_paragraph('- The generated insights, structured JSON, or images are sent back to the FastAPI layer, which returns them to the Streamlit UI to be rendered as tables, charts, chats, or visual blueprints.')

    doc.add_heading('Specialized Tool Modules', level=1)
    doc.add_heading('Filmmakers & Producers', level=2)
    doc.add_paragraph('- AI Storyboarder: Extracts scene data from BigQuery and generates visual frame references using Gemini.\n- Production Scheduler: Parses scenes, characters, props, and VFX into structured CSV-ready formats.\n- Interactive Casting: Cross-references audition videos against character profiles stored in BigQuery.')

    doc.add_heading('Writers & Script Editors', level=2)
    doc.add_paragraph('- Dynamic Script Doctor: Compares script structural beats against classic dramatic frameworks using parallel search.\n- Bespoke Dialog Partner: Lets writers chat directly with characters based on their dialogue history and subtext.\n- Localization Engine: Translates scripts while preserving regional slang and emotional tone.')

    doc.add_heading('Film Enthusiasts & Academics', level=2)
    doc.add_paragraph('- Cinematic Deep Research: Synthesizes academic-level breakdowns with precise page-level citations.\n- CYOA Simulator: Converts linear scripts into interactive role-playing games based on lore.\n- Director\'s Commentary: Maps timestamps and production notes to script chunks.')

    doc.save('SceneIQ_Architecture.docx')

if __name__ == "__main__":
    create_document()
